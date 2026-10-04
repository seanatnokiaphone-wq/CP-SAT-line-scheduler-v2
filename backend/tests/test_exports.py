"""Excel and PDF exports open, carry the expected sheets, and row counts match the plan."""
import io

import openpyxl
import pytest

from scheduler.exports import to_pdf, to_xlsx
from scheduler.generator import generate_plant
from scheduler.heuristic import best_heuristic
from scheduler.plant import FILL_LINES, SYSTEMS, default_maintenance
from scheduler.validate import validate

DT = default_maintenance()


@pytest.fixture(scope="module")
def case():
    plant = generate_plant(3, 40)
    sch = best_heuristic(plant, DT)
    return plant, sch, validate(plant, sch, DT).metrics


def _rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True)]


def test_xlsx(case, tmp_path):
    plant, sch, metrics = case
    buf = io.BytesIO()
    to_xlsx(plant, sch, DT, metrics, buf)
    wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Summary", "By line", "By tank", "Batch POs", "Fill POs"]
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.bold

    summary = {r[0]: r for r in _rows(wb["Summary"])}
    fit = summary["Batch POs finished inside the week limit"]
    assert fit[1] == metrics["pos_in_week_limit"] and fit[2] == "P0"
    assert summary["Fill POs past their hold limit"][2] == "P3"

    batches = _rows(wb["Batch POs"])
    assert len(batches) == len(plant.batches)
    assert sum(r[17] == "Yes" for r in batches) == metrics["pos_in_week_limit"]
    fills = _rows(wb["Fill POs"])
    assert len(fills) == len(plant.fills)
    assert sum((r[20] or 0) > 0 for r in fills) == metrics["fills_over_hold_limit"]
    assert {r[0] for r in fills} == {f.id for f in plant.fills}

    by_line = _rows(wb["By line"])
    assert sorted(r[7] for r in by_line if r[7]) == sorted(f.id for f in plant.fills)
    stops = [r for r in by_line if isinstance(r[5], str) and r[5].startswith("STOP")]
    assert len(stops) == len(DT)
    for line in FILL_LINES:  # each block is in time order
        hs = [r[3] for r in by_line if r[0] == line.id]
        assert hs == sorted(hs)
    by_tank = _rows(wb["By tank"])
    assert sorted(r[5] for r in by_tank if r[5]) == sorted(b.id for b in plant.batches)
    first = next(r for r in by_tank if r[5])
    assert isinstance(first[10], str) and isinstance(first[12], (int, float))  # clock string and numeric hours

    path = tmp_path / "plan.xlsx"
    to_xlsx(plant, sch, DT, metrics, str(path))
    assert openpyxl.load_workbook(path).sheetnames[0] == "Summary"


def test_pdf(case, tmp_path):
    plant, sch, metrics = case
    buf = io.BytesIO()
    to_pdf(plant, sch, DT, metrics, buf)
    data = buf.getvalue()
    assert data.startswith(b"%PDF")
    # page 1 summary + Gantt, a page per fill line, a page (or more) per system
    assert data.count(b"/Type /Page\n") + data.count(b"/Type /Page ") >= 1 + len(FILL_LINES) + len(SYSTEMS)
    path = tmp_path / "plan.pdf"
    to_pdf(plant, sch, DT, metrics, str(path))
    assert path.read_bytes().startswith(b"%PDF")
