"""Shop-floor exports of a schedule: an Excel workbook (to_xlsx) and a printable PDF pack (to_pdf).

Both take the plant, the schedule, the downtime list and the validator's metrics (validate(...).metrics), and
write to a file path or a binary buffer such as io.BytesIO. Times are given as shop-floor clock strings
("Tue 04:30") next to plain hours from Monday 07:00, so a planner can read them and a spreadsheet can sort them.
Ready, due and wait times follow validate.py (via explain._Plan), so late hours match the P3 metric.
"""
from __future__ import annotations

from .explain import _Plan, clock
from .plant import FILL_LINES, PACK_LABEL, SYSTEMS, Downtime, Plant, Schedule

TOL = 1e-4

# metric key -> (label, rule, note)
METRICS = [
    ("pos_in_week_limit", "Batch POs finished inside the week limit", "P0",
     "Batch and all its fills done by the week limit, no fill past its hold limit"),
    ("batch_pos", "Batch POs in the plan", "H13", "Every PO scheduled exactly once"),
    ("fills_over_hold_limit", "Fill POs past their hold limit", "P3", "Fill ends after batch end + hold limit"),
    ("target_miss_h", "Hours outside target windows", "P2", "Hours beyond the +/-6h window around a target start"),
    ("makespan", "Plan end (hours from Mon 07:00)", "P4", "Last batch or fill to finish"),
    ("fill_wait_h", "Fill wait, total hours", "P6", "Fill start minus the time the fill could first start"),
    ("cip_h", "CIP and wash, total hours", "P6", "Tank CIP (H5) plus line washes (H7, H10)"),
    ("tank_cip_h", "Tank CIP hours", "H5", "Minor 0.5h, standard 1.5h, deep 3h"),
    ("line_cip_h", "Line wash hours", "H7", "Includes F7 pack changes (H10)"),
]


# ---------- shared tables ----------

def _r(x, n=2):
    return None if x is None else round(float(x), n)


def _fmt(f) -> str:
    return f.format or PACK_LABEL.get(f.pack, f.pack)


def _model(plant: Plant, sch: Schedule, downtime: list[Downtime], week_limit: float = 120.0):
    """Row data shared by both exports."""
    d = _Plan(plant, sch, downtime)
    line_rows: dict[str, list[dict]] = {}
    for line in FILL_LINES:
        rows = []
        for t in d.by_line.get(line.id, []):
            f, i = d.pf[t.id], d.info[t.id]
            wash = d.wash_rule.get(t.id, "")
            rows.append({
                "kind": "fill", "line": line.id, "start": t.start, "end": t.end,
                "wash": (wash.replace("+pack", " + pack change") if wash else ""),
                "wash_h": (t.wash_end - t.wash_start) if t.wash_start is not None else 0.0,
                "wash_start": t.wash_start,
                "fill": f.id, "batch": f.batch_id, "sku": f.sku, "product": f.product_name, "cat": f.category,
                "pack": f.pack, "format": _fmt(f), "litres": f.volume_l, "units": f.units,
                "ready": i["ready"], "due": i["due"], "wait": i["wait"], "late": i["late"],
                "hold": f.hold_max})
        for z in downtime:
            if z.line == line.id:
                rows.append({"kind": "stop", "line": line.id, "start": z.start, "end": z.end, "stop": z.id,
                             "reason": f"{z.reason or 'Stop'} ({z.kind}, H14)"})
        rows.sort(key=lambda r: (r["start"], r["kind"] != "stop"))
        line_rows[line.id] = rows

    tank_rows: dict[str, list[dict]] = {}
    for sysno, tanks in SYSTEMS.items():
        for tank in tanks:
            rows = []
            for t in d.by_tank.get(tank, []):
                b = d.pb[t.id]
                fills = d.fills_of.get(d.group(b.id)[0], [])
                rows.append({
                    "tank": tank, "system": sysno, "cip": t.cip_rule or ("" if t.cip_start is None else "CIP"),
                    "cip_start": t.cip_start, "cip_end": t.cip_end,
                    "cip_h": (t.cip_end - t.cip_start) if t.cip_start is not None else 0.0,
                    "batch": b.id, "trio": b.trio_id or "", "sku": b.sku, "product": b.product_name,
                    "cat": b.category, "start": t.start, "end": t.end, "held": d.release[b.id],
                    "fills": ", ".join(f"{x} ({d.sf[x].line})" for x in fills)})
            tank_rows[tank] = rows

    batch_rows = []
    done_at = {}
    for b in plant.batches:
        grp = d.group(b.id)
        fids = d.fills_of.get(grp[0], [])
        end = max([d.sb[m].end for m in grp] + [d.sf[x].end for x in fids])
        late = any(d.info[x]["late"] > TOL for x in fids)
        done_at[b.id] = end
        t = d.sb[b.id]
        batch_rows.append({
            "batch": b.id, "trio": b.trio_id or "", "system": b.system, "tank": t.tank, "sku": b.sku,
            "product": b.product_name, "cat": b.category, "litres": b.volume_l, "duration": b.duration,
            "cip": t.cip_rule or "", "start": t.start, "end": t.end, "fills": ", ".join(fids),
            "done": end, "held": d.release[b.id],
            "in_week": end <= week_limit + TOL and not late, "late_fill": late})
    fill_rows = []
    for f in plant.fills:
        t, i = d.sf[f.id], d.info[f.id]
        fill_rows.append({
            "fill": f.id, "batch": f.batch_id, "line": t.line, "sku": f.sku, "product": f.product_name,
            "cat": f.category, "pack": f.pack, "format": _fmt(f), "litres": f.volume_l, "units": f.units,
            "duration": f.duration, "hold": f.hold_max, "ready": i["ready"], "start": t.start, "end": t.end,
            "due": i["due"], "wait": i["wait"], "late": i["late"]})
    return d, line_rows, tank_rows, batch_rows, fill_rows


def _summary_rows(sch: Schedule, metrics: dict, downtime: list[Downtime], week_limit: float):
    rows = [("Engine", sch.engine or "", "", sch.status or ""),
            ("Week limit", week_limit, "P0", clock(week_limit))]
    for key, label, rule, note in METRICS:
        if key not in metrics:
            continue
        v = metrics[key]
        if key == "makespan":
            note = f"{clock(v)}. {note}"
        rows.append((label, _r(v) if isinstance(v, float) else v, rule, note))
    for key, v in metrics.items():  # anything else the caller passed, e.g. violations by rule
        if key not in {m[0] for m in METRICS}:
            rows.append((key, v if isinstance(v, (int, float, str)) or v is None else str(v), "", ""))
    rows.append(("Planned and unplanned stops", len(downtime), "H14", "Lines are not run during a stop"))
    return rows


# ---------- Excel ----------

def to_xlsx(plant: Plant, sch: Schedule, downtime: list[Downtime], metrics: dict, path_or_buffer,
            week_limit: float = 120.0) -> None:
    """Workbook with sheets Summary, By line, By tank, Batch POs, Fill POs."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _, line_rows, tank_rows, batch_rows, fill_rows = _model(plant, sch, downtime, week_limit)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    block_font = Font(bold=True, size=12)
    wb = Workbook()

    def sheet(title, headers, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = w
        return ws

    def ck(h):
        return clock(h) if h is not None else ""

    # Summary
    ws = sheet("Summary", ["Measure", "Value", "Rule", "Notes"], [40, 14, 8, 70], first=True)
    for row in _summary_rows(sch, metrics, downtime, week_limit):
        ws.append(list(row))

    # By line
    hdr = ["Line", "Start", "End", "Start h", "End h", "Wash before", "Wash h", "Fill PO", "Batch PO", "SKU",
           "Product", "Cat", "Format", "Litres", "Units", "Due", "Due h", "Late h"]
    ws = sheet("By line", hdr, [6, 14, 14, 8, 8, 24, 7, 11, 11, 10, 30, 5, 14, 9, 9, 14, 8, 7])
    for line in FILL_LINES:
        rows = line_rows[line.id]
        packs = " / ".join(PACK_LABEL.get(p, p) for p in line.packs)
        ws.append([f"{line.id}: {packs}" + (f" (max {line.weekly_max} fills a week, H10)" if line.weekly_max else "")
                   + f", {sum(r['kind'] == 'fill' for r in rows)} fills"])
        ws.cell(ws.max_row, 1).font = block_font
        for r in rows:
            if r["kind"] == "stop":
                ws.append([r["line"], ck(r["start"]), ck(r["end"]), _r(r["start"]), _r(r["end"]),
                           f"STOP {r['stop']}: {r['reason']}"])
                ws.cell(ws.max_row, 6).font = bold
                continue
            ws.append([r["line"], ck(r["start"]), ck(r["end"]), _r(r["start"]), _r(r["end"]), r["wash"],
                       _r(r["wash_h"]), r["fill"], r["batch"], r["sku"], r["product"], r["cat"], r["format"],
                       r["litres"], r["units"], ck(r["due"]), _r(r["due"]), _r(r["late"])])
            if r["late"] > TOL:
                ws.cell(ws.max_row, 18).font = bold
        ws.append([])

    # By tank
    hdr = ["Tank", "System", "CIP", "CIP start", "CIP h", "Batch PO", "Trio", "SKU", "Product", "Cat", "Start",
           "End", "Start h", "End h", "Held until", "Held until h", "Fill POs (line)"]
    ws = sheet("By tank", hdr, [6, 7, 10, 14, 6, 11, 8, 10, 30, 5, 14, 14, 8, 8, 14, 10, 40])
    for sysno, tanks in SYSTEMS.items():
        for tank in tanks:
            rows = tank_rows[tank]
            ws.append([f"{tank} (System {sysno}), {len(rows)} batches"])
            ws.cell(ws.max_row, 1).font = block_font
            for r in rows:
                ws.append([r["tank"], r["system"], r["cip"], ck(r["cip_start"]), _r(r["cip_h"]), r["batch"],
                           r["trio"], r["sku"], r["product"], r["cat"], ck(r["start"]), ck(r["end"]),
                           _r(r["start"]), _r(r["end"]), ck(r["held"]), _r(r["held"]), r["fills"]])
            ws.append([])

    # Batch POs
    hdr = ["Batch PO", "Trio", "System", "Tank", "SKU", "Product", "Cat", "Litres", "Run h", "CIP before",
           "Start", "End", "Start h", "End h", "Fill POs", "All done", "All done h", "In week limit (P0)",
           "Fill past hold limit (P3)"]
    ws = sheet("Batch POs", hdr, [11, 8, 7, 6, 10, 30, 5, 9, 6, 10, 14, 14, 8, 8, 30, 14, 9, 10, 10])
    for r in batch_rows:
        ws.append([r["batch"], r["trio"], r["system"], r["tank"], r["sku"], r["product"], r["cat"], r["litres"],
                   r["duration"], r["cip"], ck(r["start"]), ck(r["end"]), _r(r["start"]), _r(r["end"]), r["fills"],
                   ck(r["done"]), _r(r["done"]), "Yes" if r["in_week"] else "No",
                   "Yes" if r["late_fill"] else "No"])

    # Fill POs
    hdr = ["Fill PO", "Batch PO", "Line", "SKU", "Product", "Cat", "Format", "Litres", "Units", "Fill h",
           "Hold limit h", "Ready", "Start", "End", "Due", "Ready h", "Start h", "End h", "Due h", "Wait h",
           "Late h (P3)"]
    ws = sheet("Fill POs", hdr, [11, 11, 5, 10, 30, 5, 14, 9, 9, 6, 8, 14, 14, 14, 14, 8, 8, 8, 8, 7, 8])
    for r in fill_rows:
        ws.append([r["fill"], r["batch"], r["line"], r["sku"], r["product"], r["cat"], r["format"], r["litres"],
                   r["units"], r["duration"], r["hold"], ck(r["ready"]), ck(r["start"]), ck(r["end"]), ck(r["due"]),
                   _r(r["ready"]), _r(r["start"]), _r(r["end"]), _r(r["due"]), _r(r["wait"]), _r(r["late"])])
        if r["late"] > TOL:
            ws.cell(ws.max_row, 21).font = bold
    wb.save(path_or_buffer)


# ---------- PDF ----------

CAT_GREY = {"A": 0.15, "B": 0.55, "C": 0.88}  # dark, mid, light: readable on a black-and-white printer


def _gantt(plant: Plant, sch: Schedule, downtime: list[Downtime], width: float, week_limit: float):
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib.colors import Color, black, white

    pb = {b.id: b for b in plant.batches}
    pf = {f.id: f for f in plant.fills}
    lanes = [t for ts in SYSTEMS.values() for t in ts] + [l.id for l in FILL_LINES]
    lane_h, label_w, axis_h, legend_h = 12.0, 30.0, 14.0, 14.0
    end = max([t.end for t in sch.batches] + [t.end for t in sch.fills] + [week_limit])
    span = (int(end // 12) + 1) * 12
    plot_w = width - label_w - 4
    sx = plot_w / span
    height = axis_h + len(lanes) * lane_h + legend_h + 4
    dr = Drawing(width, height)
    top = height - axis_h
    ypos = {k: top - (n + 1) * lane_h for n, k in enumerate(lanes)}
    X = lambda h: label_w + h * sx  # noqa: E731
    grey = lambda g: Color(g, g, g)  # noqa: E731

    # axis: a tick every 12h, a day label every 24h
    for h in range(0, span + 1, 12):
        dr.add(Line(X(h), top, X(h), top - len(lanes) * lane_h, strokeColor=grey(0.8), strokeWidth=0.3))
        if h % 24 == 0 and h < span:
            dr.add(String(X(h) + 1, top + 3, clock(h), fontSize=6))
    for n, k in enumerate(lanes):
        y = ypos[k]
        if k == "F1":
            dr.add(Line(label_w - 28, y + lane_h, X(span), y + lane_h, strokeColor=black, strokeWidth=0.8))
        dr.add(String(2, y + 2.5, k, fontSize=6.5))
        if n % 2:
            dr.add(Rect(label_w, y, plot_w, lane_h, fillColor=grey(0.96), strokeColor=None))

    def bar(lane, s, e, cat, label=""):
        y = ypos[lane] + 1.5
        g = CAT_GREY.get(cat, 0.7)
        dr.add(Rect(X(s), y, max(0.6, (e - s) * sx), lane_h - 3, fillColor=grey(g), strokeColor=black,
                    strokeWidth=0.3))
        if label and (e - s) * sx > 5.5:
            dr.add(String(X(s) + 1.2, y + 1.5, label, fontSize=5, fillColor=white if g < 0.5 else black))

    def hatch(lane, s, e):
        y0 = ypos[lane] + 1
        hgt = lane_h - 2
        x0, x1 = X(s), X(e)
        dr.add(Rect(x0, y0, x1 - x0, hgt, fillColor=white, strokeColor=black, strokeWidth=0.4))
        x = x0
        while x < x1:
            dr.add(Line(x, y0, min(x + hgt, x1), y0 + min(hgt, x1 - x), strokeColor=black, strokeWidth=0.3))
            x += 2.5

    for t in sch.batches:
        if t.cip_start is not None and t.cip_end > t.cip_start:
            dr.add(Rect(X(t.cip_start), ypos[t.tank] + 3.5, (t.cip_end - t.cip_start) * sx, lane_h - 7,
                        fillColor=white, strokeColor=black, strokeWidth=0.3))
        b = pb[t.id]
        bar(t.tank, t.start, t.end, b.category, b.category)
    for t in sch.fills:
        if t.wash_start is not None and t.wash_end > t.wash_start:
            dr.add(Rect(X(t.wash_start), ypos[t.line] + 3.5, (t.wash_end - t.wash_start) * sx, lane_h - 7,
                        fillColor=white, strokeColor=black, strokeWidth=0.3))
        f = pf[t.id]
        bar(t.line, t.start, t.end, f.category, f.category)
    for z in downtime:
        if z.line in ypos:
            hatch(z.line, z.start, z.end)

    # week limit
    wx = X(week_limit)
    dr.add(Line(wx, top + 1, wx, top - len(lanes) * lane_h, strokeColor=black, strokeWidth=1.2,
                strokeDashArray=[3, 2]))
    dr.add(String(wx + 2, top - len(lanes) * lane_h - 8, f"Week limit {clock(week_limit)} ({week_limit:g}h, P0)",
                  fontSize=6))

    # legend
    lx, ly = label_w, 2
    for cat in ("A", "B", "C"):
        dr.add(Rect(lx, ly, 14, 7, fillColor=grey(CAT_GREY[cat]), strokeColor=black, strokeWidth=0.3))
        dr.add(String(lx + 17, ly + 1, f"Category {cat}", fontSize=6))
        lx += 70
    dr.add(Rect(lx, ly + 2, 14, 3, fillColor=white, strokeColor=black, strokeWidth=0.3))
    dr.add(String(lx + 17, ly + 1, "CIP / wash (H5, H7)", fontSize=6))
    lx += 85
    dr.add(Rect(lx, ly, 14, 7, fillColor=white, strokeColor=black, strokeWidth=0.4))
    for k in range(0, 14, 3):
        dr.add(Line(lx + k, ly, lx + min(k + 7, 14), ly + min(7, 14 - k), strokeColor=black, strokeWidth=0.3))
    dr.add(String(lx + 17, ly + 1, "Planned stop (H14)", fontSize=6))
    return dr


def to_pdf(plant: Plant, sch: Schedule, downtime: list[Downtime], metrics: dict, path_or_buffer,
           week_limit: float = 120.0) -> None:
    """Landscape A4 pack: summary and Gantt on page 1, then one table per fill line and per system."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    _, line_rows, tank_rows, _, _ = _model(plant, sch, downtime, week_limit)
    page = landscape(A4)
    margin = 10 * mm
    width = page[0] - 2 * margin
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=14, spaceAfter=4, spaceBefore=0)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, spaceAfter=3, spaceBefore=0)
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=7, leading=8.5)
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=7, leading=8)
    end = metrics.get("makespan", max([t.end for t in sch.batches] + [t.end for t in sch.fills], default=0.0))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(margin, 6 * mm, f"Plant seed {plant.seed} · {sch.engine} plan · times from Mon 07:00")
        canvas.drawRightString(page[0] - margin, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    grid = TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.85, 0.85, 0.85)),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)])

    def table(header, rows, widths, extra=()):
        scale = width / sum(widths)
        t = LongTable([header] + rows, colWidths=[w * scale for w in widths], repeatRows=1)
        t.setStyle(grid)
        if extra:
            t.setStyle(TableStyle(list(extra)))
        return t

    story = [Paragraph(f"Production plan: {len(plant.batches)} batch POs, {len(plant.fills)} fill POs", h1)]
    fit = metrics.get("pos_in_week_limit")
    late = metrics.get("fills_over_hold_limit")
    story.append(Paragraph(
        f"{fit if fit is not None else '?'} of {len(plant.batches)} batch POs finish inside the week limit of "
        f"{clock(week_limit)} (P0); {late if late is not None else '?'} fill POs run past their hold limit (P3). "
        f"The plan ends {clock(end)} ({end:g}h).", small))
    story.append(Spacer(1, 3))
    srows = _summary_rows(sch, metrics, downtime, week_limit)
    half = (len(srows) + 1) // 2
    left, right = srows[:half], srows[half:] + [("", "", "", "")] * (2 * half - len(srows))
    body = [[a[0], str(a[1]), a[2], b[0], str(b[1]), b[2]] for a, b in zip(left, right)]
    st = Table([["Measure", "Value", "Rule", "Measure", "Value", "Rule"]] + body,
               colWidths=[w * width / 100 for w in (34, 10, 6, 34, 10, 6)])
    st.setStyle(grid)
    story += [st, Spacer(1, 5), _gantt(plant, sch, downtime, width, week_limit), PageBreak()]

    # one table per fill line
    hdr = ["Start", "End", "Wash before", "Fill PO", "Batch PO", "SKU", "Product", "Cat", "Format", "Litres",
           "Units", "Due", "Late", "Done"]
    widths = [14, 14, 16, 11, 11, 9, 28, 4, 13, 8, 7, 14, 7, 6]
    for k, line in enumerate(FILL_LINES):
        rows, extra = [], []
        packs = " / ".join(PACK_LABEL.get(p, p) for p in line.packs)
        for r in line_rows[line.id]:
            n = len(rows) + 1
            if r["kind"] == "stop":
                rows.append([clock(r["start"]), clock(r["end"]), "STOP " + r["stop"], r["reason"]] + [""] * 10)
                extra += [("SPAN", (3, n), (-1, n)), ("FONT", (0, n), (-1, n), "Helvetica-Bold", 7)]
                continue
            wash = f"{r['wash']} {r['wash_h']:g}h" if r["wash"] else ""
            lt = f"+{r['late']:.1f}h" if r["late"] > TOL else ""
            rows.append([clock(r["start"]), clock(r["end"]), wash, r["fill"], r["batch"], r["sku"],
                         Paragraph(r["product"], cell), r["cat"], r["format"], r["litres"] or "", r["units"] or "",
                         clock(r["due"]), lt, "[  ]"])
            if lt:
                extra.append(("FONT", (12, n), (12, n), "Helvetica-Bold", 7))
        nf = sum(x["kind"] == "fill" for x in line_rows[line.id])
        story.append(Paragraph(f"Fill line {line.id}: {packs}, {nf} fills"
                               + (f" (max {line.weekly_max} a week, H10)" if line.weekly_max else ""), h2))
        story.append(table(hdr, rows or [["No fills"] + [""] * 13], widths, extra))
        story.append(PageBreak())

    # one table per system
    hdr = ["Tank", "CIP before", "CIP start", "Batch PO", "Trio", "SKU", "Product", "Cat", "Start", "End",
           "Held until", "Fill POs (line)", "Done"]
    widths = [6, 10, 14, 11, 8, 9, 26, 4, 14, 14, 14, 30, 6]
    for sysno, tanks in SYSTEMS.items():
        rows = []
        for r in sorted((r for tk in tanks for r in tank_rows[tk]), key=lambda r: (r["start"], r["tank"])):
            cip = f"{r['cip']} {r['cip_h']:g}h" if r["cip_start"] is not None else ""
            rows.append([r["tank"], cip, clock(r["cip_start"]) if r["cip_start"] is not None else "", r["batch"],
                         r["trio"], r["sku"], Paragraph(r["product"], cell), r["cat"], clock(r["start"]),
                         clock(r["end"]), clock(r["held"]), Paragraph(r["fills"], cell), "[  ]"])
        story.append(Paragraph(f"System {sysno}: tanks {', '.join(tanks)}, {len(rows)} batches", h2))
        story.append(table(hdr, rows or [["No batches"] + [""] * 12], widths))
        if sysno != max(SYSTEMS):
            story.append(PageBreak())

    doc = SimpleDocTemplate(path_or_buffer, pagesize=page, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=12 * mm, title="Production plan",
                            author="cp-sat-line-scheduler")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


__all__ = ["to_xlsx", "to_pdf"]
